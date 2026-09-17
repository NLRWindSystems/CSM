"""Generates an NLR COWE Review-style CapEx breakdown table for one offshore wind (OSW) turbine
configuration, formatted exactly like
:py:mod:`csm.tools.experimental.nlr_cower_table_generator` (the land-based table this mirrors) —
same indented, hierarchical "Parameter" column using sentence case throughout, same $/kW values
(cost divided by that configuration's ``rated_power_kw``), same workbook styling. Unlike the
land-based table's one-column-per-model-year layout, each row here carries one column per
offshore type (:py:data:`OSW_MODELS`) — Fixed-Bottom and Floating — so the single workbook
doubles as a fixed-bottom-vs-floating comparison for that one configuration.

The line-item -> OSWBase cost field mapping (:py:data:`COWER_LINE_ITEMS`) follows the breakdown
requested for this table specifically: Turbine > {Rotor-nacelle assembly (Rotor + Nacelle),
Tower} — narrower than OSWBase's own four top-level assemblies (Blade isn't a line item here,
since it's already counted inside Rotor; see :py:mod:`csm.models.base_osw_model`), and not
reusing any of build_custom_osw_model.py's own MASS_SPECS/COST_SPECS groupings — every row here
sums OSWBase's own per-assembly cost attributes directly.
"""

from typing import NamedTuple
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Side, Border, Alignment, PatternFill

from csm.models.nlr2024_osw_fb import OffshoreFB2024NLR
from csm.models.nlr2024_osw_fl import OffshoreFL2024NLR


# Deliberately duplicated from osw_plotting.py rather than imported from it — this script is
# meant to run as its own independent entry point, not to depend on (or be broken by unrelated
# changes to) osw_plotting.py's own module state. Keep these in sync by hand if either changes.
OSW_MODELS: dict[str, type] = {
    "Fixed-Bottom": OffshoreFB2024NLR,
    "Floating": OffshoreFL2024NLR,
}

# The single configuration this table is scoped to: a 12 MW turbine with a 216 m rotor diameter
# and a 137 m hub height.
DEFAULT_TURBINE_SPECS: dict[str, dict] = {
    "12 MW (216 m RD, 137 m HH)": {
        "rated_power_MW": 12.0,
        "rotor_diameter": 216.0,
        "hub_height": 137.0,
    },
}


def to_model_kwargs(spec: dict) -> dict:
    """Converts a raw turbine specification into ``OSWBase`` subclass constructor kwargs.

    Args:
        spec (dict): Turbine specification with keys ``rated_power_MW``, ``rotor_diameter``, and
            ``hub_height``.

    Returns:
        dict: Keyword arguments accepted by every model in :py:data:`OSW_MODELS`.
    """
    return {
        "rated_power_kw": round(spec["rated_power_MW"] * 1000),
        "rotor_diameter": float(spec["rotor_diameter"]),
        "hub_height": float(spec["hub_height"]),
    }


class LineItem(NamedTuple):
    """One row of the COWER-style table.

    Either `cost_attrs` (a leaf row, summing real OSWBase cost attributes) or `subtotal_of` (a
    rollup row, summing other line items' *already-computed* $/kW values by label) is set, never
    both — a rollup is always exactly consistent with its own children instead of being
    independently recomputed from cost attributes of its own.

    Each `cost_attrs` entry is `(attr_name, multiplier_key)` — `multiplier_key` is always `None`
    here (unlike the land-based table, OSWBase has no per-blade/per-bearing style attributes that
    need scaling before summing), kept only so this mirrors the land-based table's `LineItem`
    shape exactly.
    """

    label: str
    indent: int
    cost_attrs: tuple[tuple[str, str | None], ...] = ()
    subtotal_of: tuple[str, ...] = ()
    bold: bool = False


# The breakdown given for this table: Turbine > {Rotor-nacelle assembly (Rotor + Nacelle), Tower}.
COWER_LINE_ITEMS: list[LineItem] = [
    LineItem("Turbine", 0, subtotal_of=("Rotor-nacelle assembly", "Tower"), bold=True),
    LineItem("Rotor-nacelle assembly", 1, subtotal_of=("Rotor", "Nacelle"), bold=True),
    LineItem("Rotor", 2, cost_attrs=(("rotor_cost", None),)),
    LineItem("Nacelle", 2, cost_attrs=(("nacelle_cost", None),)),
    LineItem("Tower", 1, cost_attrs=(("tower_cost", None),), bold=True),
]

_ITEMS_BY_LABEL: dict[str, LineItem] = {item.label: item for item in COWER_LINE_ITEMS}

HEADER_FILL = PatternFill("solid", fgColor="000000")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ACCENT_FILL = PatternFill("solid", fgColor="2E74B5")
THIN_SIDE = Side(style="thin", color="BFBFBF")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
VALUE_NUMBER_FORMAT = "#,##0"


def _row_value_per_kw(
    label: str, model, rated_power_kw: float, model_kwargs: dict, cache: dict[str, float]
) -> float:
    """The $/kW value of line item `label` for one already-run `model` instance, memoized in
    `cache` so a rollup row (e.g. "Turbine") only walks its children's values once each,
    regardless of how many ancestor rollups also depend on them.
    """
    if label in cache:
        return cache[label]
    item = _ITEMS_BY_LABEL[label]
    if item.subtotal_of:
        value = sum(
            _row_value_per_kw(child, model, rated_power_kw, model_kwargs, cache)
            for child in item.subtotal_of
        )
    else:
        raw_cost = sum(
            (getattr(model, attr, 0.0) or 0.0)
            * (model_kwargs[multiplier_key] if multiplier_key is not None else 1)
            for attr, multiplier_key in item.cost_attrs
        )
        value = raw_cost / rated_power_kw
    cache[label] = value
    return value


def _compute_config_values(
    spec: dict, models: dict[str, type]
) -> dict[str, dict[str, float | None]]:
    """Runs every model against one turbine configuration and computes every line item's $/kW
    value for each.

    Returns:
        dict[str, dict[str, float | None]]: label -> model name -> $/kW value, or None for a
            model/configuration combination that raised while running (a model rejecting an
            infeasible combination of inputs, say) rather than letting the whole workbook fail.
    """
    model_kwargs = to_model_kwargs(spec)
    rated_power_kw = model_kwargs["rated_power_kw"]
    values: dict[str, dict[str, float | None]] = {item.label: {} for item in COWER_LINE_ITEMS}
    for model_name, model_cls in models.items():
        try:
            model = model_cls(**model_kwargs)
            model.run()
        except (ValueError, AttributeError, TypeError):
            for item in COWER_LINE_ITEMS:
                values[item.label][model_name] = None
            continue
        cache: dict[str, float] = {}
        for item in COWER_LINE_ITEMS:
            values[item.label][model_name] = _row_value_per_kw(
                item.label, model, rated_power_kw, model_kwargs, cache
            )
    return values


def _slugify_filename(name: str) -> str:
    keep = "".join(c if c.isalnum() or c in "()._-" else "_" for c in name)
    while "__" in keep:
        keep = keep.replace("__", "_")
    return keep.strip("_")


def _write_workbook(
    config_name: str,
    model_names: list[str],
    values: dict[str, dict[str, float | None]],
    output_path: Path,
) -> None:
    """Writes one config's table (see module docstring for the layout) to `output_path`."""
    wb = Workbook()
    ws = wb.active
    ws.title = "CapEx Breakdown"

    n_cols = 1 + len(model_names)
    ws.cell(row=1, column=1, value="Parameter")
    for col, model_name in enumerate(model_names, start=2):
        ws.cell(row=1, column=col, value=f"{model_name} ($/kW)")
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for col in range(1, n_cols + 1):
        ws.cell(row=2, column=col).fill = ACCENT_FILL
    ws.row_dimensions[2].height = 6

    for row_offset, item in enumerate(COWER_LINE_ITEMS):
        row = row_offset + 3
        label_cell = ws.cell(row=row, column=1, value=item.label)
        label_cell.alignment = Alignment(indent=item.indent, horizontal="left", vertical="center")
        label_cell.font = Font(bold=item.bold)
        label_cell.border = THIN_BORDER
        for col, model_name in enumerate(model_names, start=2):
            value = values[item.label][model_name]
            value_cell = ws.cell(row=row, column=col, value=value if value is not None else "n/a")
            if value is not None:
                value_cell.number_format = VALUE_NUMBER_FORMAT
            value_cell.font = Font(bold=item.bold)
            value_cell.alignment = Alignment(horizontal="right", vertical="center")
            value_cell.border = THIN_BORDER

    ws.column_dimensions["A"].width = 34
    for col in range(2, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.freeze_panes = "B3"
    ws.print_title_rows = "1:1"

    wb.properties.title = f"NLR COWE Review-style CapEx breakdown (OSW) — {config_name}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))


def generate_cower_osw_tables(
    output_dir: str | Path = "output/cower_osw_tables",
    models: dict[str, type] | None = None,
    configs: dict[str, dict] | None = None,
) -> dict[str, Path]:
    """Generates one NLR COWER-style CapEx breakdown workbook per OSW turbine configuration.

    Args:
        output_dir (str | Path, optional): Directory the workbooks are saved into (created if
            missing). Defaults to "output/cower_osw_tables".
        models (dict[str, type] | None, optional): Mapping of offshore type to an ``OSWBase``
            subclass — becomes one $/kW column per row. Defaults to :py:data:`OSW_MODELS`.
        configs (dict[str, dict] | None, optional): Mapping of configuration name to a raw
            turbine spec (see :py:func:`to_model_kwargs`) — becomes one workbook per entry.
            Defaults to :py:data:`DEFAULT_TURBINE_SPECS`.

    Returns:
        dict[str, Path]: Mapping of configuration name to its saved workbook path.
    """
    models = models or OSW_MODELS
    configs = configs or DEFAULT_TURBINE_SPECS
    output_dir = Path(output_dir).resolve()

    model_names = list(models)
    paths: dict[str, Path] = {}
    for config_name, spec in configs.items():
        values = _compute_config_values(spec, models)
        path = output_dir / f"cower_osw_{_slugify_filename(config_name)}.xlsx"
        _write_workbook(config_name, model_names, values, path)
        paths[config_name] = path
    return paths


if __name__ == "__main__":
    for config_name, path in generate_cower_osw_tables().items():
        print(f"{config_name}: {path}")  # noqa: T201 — CLI entry point, not a debug leftover
